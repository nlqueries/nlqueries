"""Bounds on how much work extracting one document may do (finding 2.8).

The companion to ``test_document_expansion_limits.py``. That gate reads what an
archive *declares* about itself before parsing; these bound what extraction
actually does while reading, so they hold for a `.pdf` (not an archive at all)
and for an archive whose directory misstated its contents.

Real files and real parsers, for the same reason as the expansion tests: a
mocked ``openpyxl`` would confirm the arithmetic and say nothing about whether a
real workbook is refused.
"""

from __future__ import annotations

import time
from pathlib import Path
from typing import Any

import pytest
from nlqueries.document_connectors import _limits
from nlqueries.document_connectors._limits import (
    DocumentExtractionTimeout,
    DocumentTooComplexError,
    ExtractionBudget,
    check_archive_expansion,
)

openpyxl = pytest.importorskip("openpyxl")
docx = pytest.importorskip("docx")

_BATCH = 50


def _sheet(path: Path, rows: int, *, header: bool = True, width: int = 3) -> Path:
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("data")
    if header:
        ws.append([f"col{i}" for i in range(width)])
    for n in range(rows):
        ws.append([f"r{n}"] * width)
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------------
# The row cap
# ---------------------------------------------------------------------------


def test_a_workbook_under_the_row_cap_is_ingested(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The control. Every limit here is a ceiling over documents people really have."""
    from nlqueries.document_connectors.excel import ExcelConnector

    monkeypatch.setattr(_limits, "MAX_ROWS", 200)
    chunks = ExcelConnector().ingest(_sheet(tmp_path / "ok.xlsx", rows=200), source_id="s")
    assert len(chunks) == 4, [c.chunk_index for c in chunks]


def test_a_workbook_over_the_row_cap_is_refused(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Counted while reading, so it bounds an archive that lied about its size."""
    from nlqueries.document_connectors.excel import ExcelConnector

    monkeypatch.setattr(_limits, "MAX_ROWS", 200)
    with pytest.raises(DocumentTooComplexError, match="more than 200 rows"):
        ExcelConnector().ingest(_sheet(tmp_path / "big.xlsx", rows=201), source_id="s")


def test_the_row_cap_counts_across_sheets_not_per_sheet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A workbook is refused on its total, not on its largest sheet.

    Per-sheet would be trivially evaded by splitting the same rows across
    several tabs, which costs a worker exactly as much.
    """
    from nlqueries.document_connectors.excel import ExcelConnector

    path = tmp_path / "many.xlsx"
    wb = openpyxl.Workbook(write_only=True)
    for s in range(4):
        ws = wb.create_sheet(f"s{s}")
        ws.append(["a", "b"])
        for n in range(60):
            ws.append([f"v{n}", n])
    wb.save(str(path))

    monkeypatch.setattr(_limits, "MAX_ROWS", 200)  # 4 x 60 = 240 total, 60 per sheet
    with pytest.raises(DocumentTooComplexError, match="more than 200 rows"):
        ExcelConnector().ingest(path, source_id="s")


# ---------------------------------------------------------------------------
# The refactor that made the row cap possible: streaming, not accumulating
# ---------------------------------------------------------------------------


def test_a_blank_batch_still_advances_the_chunk_index(tmp_path: Path) -> None:
    """The subtlety the streaming rewrite had to preserve exactly.

    The old connector built every row of a sheet, enumerated over fixed offsets
    and skipped a batch whose text was blank -- so a blank stretch leaves a *gap*
    in the emitted indexes rather than renumbering what follows. Chunk ids are
    derived from that index, so getting it wrong silently re-ids every chunk
    after a blank run, and a re-ingest would orphan the old vectors instead of
    replacing them.

    Measured against the pre-refactor connector: indexes ``[0, 2]`` and row
    ranges ``2-51`` / ``102-121``.
    """
    from nlqueries.document_connectors.excel import ExcelConnector

    path = tmp_path / "gap.xlsx"
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("data")
    ws.append(["alpha", "beta", "gamma"])
    for n in range(1, 51):  # batch 0: 50 rows
        ws.append([f"r{n}", n, n * 1.5])
    for _ in range(50):  # batch 1: blank, skipped but still counted
        ws.append([None, None, None])
    for n in range(1, 21):  # batch 2: 20 rows
        ws.append([f"r{n}", n, n * 1.5])
    wb.save(str(path))

    chunks = ExcelConnector().ingest(path, source_id="fixed-source-id")

    assert [c.chunk_index for c in chunks] == [0, 2]
    assert [(c.metadata or {})["row_range"] for c in chunks] == ["2-51", "102-121"]


def test_the_last_partial_batch_is_emitted(tmp_path: Path) -> None:
    """Streaming emits on a full batch, so the remainder needs its own flush.

    Forgetting it is the obvious way to break this refactor, and it would drop
    the tail of every sheet whose row count is not a multiple of 50 -- quietly,
    since the chunks that remain are all correct.
    """
    from nlqueries.document_connectors.excel import ExcelConnector

    chunks = ExcelConnector().ingest(_sheet(tmp_path / "tail.xlsx", rows=125), source_id="s")

    assert [c.chunk_index for c in chunks] == [0, 1, 2]
    assert (chunks[-1].metadata or {})["row_range"] == "102-126"


def test_a_header_only_sheet_produces_nothing(tmp_path: Path) -> None:
    """No data rows means no partial batch to flush, and no empty chunk either."""
    from nlqueries.document_connectors.excel import ExcelConnector

    assert ExcelConnector().ingest(_sheet(tmp_path / "h.xlsx", rows=0), source_id="s") == []


# ---------------------------------------------------------------------------
# The wall-clock budget
# ---------------------------------------------------------------------------


def test_the_budget_allows_an_ordinary_document() -> None:
    """The control for the clock: a budget that has not elapsed must not fire."""
    ExtractionBudget(seconds=60, name="x").check("0 pages")  # must not raise


def test_the_budget_refuses_once_it_is_spent() -> None:
    """And names how far extraction got, which is what makes the error useful."""
    budget = ExtractionBudget(seconds=0.01, name="slow.pdf")
    time.sleep(0.05)
    with pytest.raises(DocumentTooComplexError, match=r"slow\.pdf.*budget.*412 of 5000 pages"):
        budget.check("412 of 5000 pages")


def test_the_budget_defaults_to_the_configured_ceiling(monkeypatch: pytest.MonkeyPatch) -> None:
    """Read at call time, so a deployment's setting and a test both take effect.

    A small non-zero ceiling and a sleep well past it, rather than `0.0` and a
    strict `>`. On Windows under Python 3.11 and 3.12 `time.monotonic()` is
    backed by `GetTickCount64()` at roughly 15.6 ms, so a `0.0` budget checked
    within one tick reports `elapsed` of exactly `0.0` and raises nothing. CI
    runs Linux and would never have seen it; the Windows checkout would.
    """
    monkeypatch.setattr(_limits, "MAX_SECONDS", 0.01)
    budget = ExtractionBudget(name="x")
    time.sleep(0.05)
    with pytest.raises(DocumentTooComplexError):
        budget.check("somewhere")


@pytest.mark.parametrize("connector_name", ["excel", "word", "pdf"])
def test_every_connector_is_bounded_by_the_clock(
    connector_name: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """One budget across all three, rather than a count cap invented per format.

    Parameterised because the point of a shared budget is that no connector is
    exempt -- and `pdf` is the one that matters most, being the only format the
    expansion gate never sees (a PDF is not a zip).

    Each parser is made slow and the ceiling set well under it. An earlier
    revision patched the ceiling to `0.0` instead and relied on the clock having
    advanced at all by the first check. On Windows under Python 3.11 and 3.12
    `time.monotonic()` is backed by `GetTickCount64()` at roughly 15.6 ms, which
    these small fixtures finish well inside, so `elapsed` was exactly `0.0`, the
    strict `>` refused nothing and the test failed -- on a Windows checkout only.
    CI runs Linux and would never have shown it.

    This asserts that each connector is bounded at all. *Where* each one starts
    its clock is a separate claim, asserted below by the stage the message names.
    """
    monkeypatch.setattr(_limits, "MAX_SECONDS", 0.05)

    def _slow(real: Any) -> Any:
        """Stand in for a document that is genuinely expensive to parse.

        0.2s against a 0.05s ceiling: four times the budget and an order of
        magnitude past the coarsest clock either platform has, so nothing here
        turns on timing resolution.
        """

        def _wrapper(*args: Any, **kwargs: Any) -> Any:
            time.sleep(0.2)
            return real(*args, **kwargs)

        return _wrapper

    # Patched on each library module rather than on the connector module: all
    # three import their parser *inside* `ingest`, so there is no attribute on
    # the connector to replace and the name resolves at call time.
    if connector_name == "excel":
        from nlqueries.document_connectors.excel import ExcelConnector

        connector: Any = ExcelConnector()
        path = _sheet(tmp_path / "x.xlsx", rows=120)
        monkeypatch.setattr(openpyxl, "load_workbook", _slow(openpyxl.load_workbook))
    elif connector_name == "word":
        from nlqueries.document_connectors.word import WordConnector

        connector = WordConnector()
        path = tmp_path / "x.docx"
        document = docx.Document()
        for n in range(40):
            document.add_paragraph(f"paragraph {n} " * 10)
        document.save(str(path))
        monkeypatch.setattr(docx, "Document", _slow(docx.Document))
    else:
        pdfplumber = pytest.importorskip("pdfplumber")
        from nlqueries.document_connectors.pdf import PdfConnector

        connector = PdfConnector()
        path = _minimal_pdf(tmp_path / "x.pdf", pages=3)
        monkeypatch.setattr(pdfplumber, "open", _slow(pdfplumber.open))

    with pytest.raises(DocumentTooComplexError, match="budget"):
        connector.ingest(path, source_id="s")


def _minimal_pdf(path: Path, pages: int) -> Path:
    """A syntactically valid multi-page PDF, written by hand.

    No PDF *writer* is among the dev dependencies -- pdfplumber only reads -- and
    adding one to test a clock would be a dependency for a stopwatch.
    """
    objects: list[bytes] = []
    kids = " ".join(f"{3 + i} 0 R" for i in range(pages))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {pages} >>".encode())
    for _ in range(pages):
        objects.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
    out += trailer.encode() + f"startxref\n{xref_at}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def test_the_word_budget_covers_parsing_not_only_chunking(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Where the clock starts, which the parameterised test above cannot see.

    That test matches the word "budget" and nothing else. Measured, by deleting
    `budget.check("parsing the document")` — it still passes, reporting "0 of 1
    sections" from the check further down; this one fails. Matching the stage is
    what holds parsing inside the clock.
    """
    from nlqueries.document_connectors import word as word_module

    path = tmp_path / "slow.docx"
    document = docx.Document()
    for n in range(5):
        document.add_paragraph(f"paragraph {n}")
    document.save(str(path))

    real_document = docx.Document

    def _slow_document(*args: Any, **kwargs: Any) -> Any:
        time.sleep(0.2)  # stand in for a document that is expensive to parse
        return real_document(*args, **kwargs)

    # Patched on the `docx` module itself, not on `word_module`: the connector
    # imports docx *inside* `ingest`, so there is no module attribute to replace
    # and it resolves the name from `sys.modules` at call time.
    monkeypatch.setattr(docx, "Document", _slow_document)
    monkeypatch.setattr(_limits, "MAX_SECONDS", 0.05)

    with pytest.raises(DocumentTooComplexError, match="parsing the document"):
        word_module.WordConnector().ingest(path, source_id="s")


def test_padding_rows_do_not_count_towards_the_row_cap(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The cap must count rows a person can count in their own file.

    `iter_rows` yields every *materialised* row, and whole-column formatting
    leaves thousands of empty ones behind. Measured on this version: a 200-row
    sheet with such padding yields 5,201 tuples. A cap counting those refuses a
    workbook its owner sees 200 rows in.

    (The reviewer proposed the mechanism as `<dimension>` over-declaration.
    Measured, that does *not* pad — a sheet declaring `A1:C1048576` still yields
    only its real rows. Materialised blank rows are the route that does.)
    """
    from nlqueries.document_connectors.excel import ExcelConnector

    path = tmp_path / "padded.xlsx"
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("data")
    ws.append(["a", "b", "c"])
    for n in range(100):
        ws.append([f"r{n}", n, n * 2])
    for _ in range(1_000):
        ws.append([None, None, None])
    wb.save(str(path))

    monkeypatch.setattr(_limits, "MAX_ROWS", 200)  # 100 real rows, 1,100 yielded
    chunks = ExcelConnector().ingest(path, source_id="s")  # must not raise
    assert chunks, "a padded workbook produced no chunks"


def test_the_seconds_setting_falls_back_rather_than_raising(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A malformed value must not raise while `nlqueries.config` is imported.

    That would take down the CLI and the MCP server at startup, not merely fail
    one ingest -- and `docs/configuration.md` promises fallback behaviour for the
    neighbouring document limits, so an operator expects it here.

    `nan` has its own case: every comparison against it is false, so a budget of
    `nan` would not fail loudly, it would silently never expire.
    """
    from nlqueries.config import _positive_float

    monkeypatch.setenv("NLQ_MAX_EXTRACTION_SECONDS", "45.5")
    assert _positive_float("NLQ_MAX_EXTRACTION_SECONDS", 120.0) == 45.5

    for bad in ("120s", "", "abc", "0", "-1", "inf", "-inf", "nan"):
        monkeypatch.setenv("NLQ_MAX_EXTRACTION_SECONDS", bad)
        assert _positive_float("NLQ_MAX_EXTRACTION_SECONDS", 120.0) == 120.0, bad


def test_the_excel_budget_covers_opening_the_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Where the Excel clock starts, for the same reason as the Word case.

    `test_every_connector_is_bounded_by_the_clock[excel]` asserts only that a
    refusal happens: it matches the word "budget" and nothing else. Measured, by
    deleting `budget.check("opening the workbook")` and leaving the budget where
    it is — that test still passes, because the next check inside the sheet loop
    absorbs the refusal and reports "sheet 1"; this one fails. Opening is covered
    only because something asserts the *stage*.

    `read_only=True` makes row iteration lazy, not opening: measured at ~0.33s
    for a 60,000-row workbook before a single row is read, so opening is a real
    cost to leave outside the clock.
    """
    import openpyxl as real_openpyxl
    from nlqueries.document_connectors.excel import ExcelConnector

    path = _sheet(tmp_path / "slow.xlsx", rows=10)
    real_load = real_openpyxl.load_workbook

    def _slow_load(*args: Any, **kwargs: Any) -> Any:
        time.sleep(0.2)
        return real_load(*args, **kwargs)

    # Patched on `openpyxl` itself: the connector imports it inside `ingest`, so
    # there is no attribute on the connector module to replace.
    monkeypatch.setattr(real_openpyxl, "load_workbook", _slow_load)
    monkeypatch.setattr(_limits, "MAX_SECONDS", 0.05)

    with pytest.raises(DocumentTooComplexError, match="opening the workbook"):
        ExcelConnector().ingest(path, source_id="s")


def test_only_the_clock_raises_the_retryable_refusal(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The split that decides whether a caller may retry.

    A refusal is worth retrying exactly when it is not a property of the file.
    The expansion gate reads the archive's own directory and the row cap counts
    rows: retry either and you re-download and re-parse to reach the identical
    answer. The budget measures wall-clock, so it is a function of what else the
    machine was doing, and the same document can fail it once and pass next time.

    Asserted in both directions. The positive case alone is satisfied by making
    *everything* a timeout, which turns a genuine refusal into an unbounded retry
    loop -- the opposite defect, and the more expensive one.
    """
    budget = ExtractionBudget(seconds=0.01, name="slow.xlsx")
    time.sleep(0.05)
    with pytest.raises(DocumentExtractionTimeout):
        budget.check("somewhere")

    # ...and still catchable as the base class, so every handler that existed
    # before this subclass keeps working across the ref bump introducing it.
    spent = ExtractionBudget(seconds=0.01, name="slow.xlsx")
    time.sleep(0.05)
    with pytest.raises(DocumentTooComplexError):
        spent.check("somewhere")

    # Control: the deterministic gates are not timeouts.
    monkeypatch.setattr(_limits, "MAX_EXPANDED_BYTES", 1)
    with pytest.raises(DocumentTooComplexError) as expansion:
        check_archive_expansion(_sheet(tmp_path / "big.xlsx", rows=50))
    assert not isinstance(expansion.value, DocumentExtractionTimeout), (
        "the expansion gate reads a declared size; a retry pays for a download "
        "and a parse to be told the same thing"
    )

    monkeypatch.setattr(_limits, "MAX_EXPANDED_BYTES", 10 * 1024**3)
    monkeypatch.setattr(_limits, "MAX_EXPANSION_RATIO", 10_000)
    monkeypatch.setattr(_limits, "MAX_ROWS", 10)
    from nlqueries.document_connectors.excel import ExcelConnector

    with pytest.raises(DocumentTooComplexError) as rows:
        ExcelConnector().ingest(_sheet(tmp_path / "wide.xlsx", rows=200), source_id="s")
    assert not isinstance(rows.value, DocumentExtractionTimeout), (
        "the row cap counts rows in the file; the count is the same on a retry"
    )


def test_the_package_exports_the_refusals_it_expects_callers_to_name() -> None:
    """`_limits` is private; the exceptions in it are not.

    Enterprise's ingestion task names both, to tell a refusal from a failure and
    a retryable refusal from a final one -- and it is pinned to a core *commit*
    rather than a release, so an import reaching into `_limits` would break on a
    rename core is entitled to make without warning. `__all__` is the declared
    surface that makes that dependency legitimate; this fails if either name
    leaves it, or if the package rebinds one to something else.
    """
    import nlqueries.document_connectors as package

    for name in ("DocumentTooComplexError", "DocumentExtractionTimeout"):
        assert name in package.__all__, f"{name} left the package's declared surface"
        assert getattr(package, name) is getattr(_limits, name)


def test_the_pdf_budget_covers_opening_the_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Where the PDF clock starts -- the third of the three, added by review.

    The PDF connector has no `check()` between constructing the budget and
    `pdfplumber.open`; the first one is inside the page loop, so opening is
    covered only by the budget being built above the `with`. PDFs are the one
    format the expansion gate never sees, so the clock is all that bounds them.

    The assertion is the stage rather than the word "budget": "0 of 3 pages" can
    only be produced before any page was extracted. Moving the construction into
    the `with` fails this and the parameterised case both; deleting the first
    check instead would leave only this kind of assertion to notice.
    """
    pdfplumber = pytest.importorskip("pdfplumber")
    from nlqueries.document_connectors.pdf import PdfConnector

    path = _minimal_pdf(tmp_path / "slow.pdf", pages=3)
    real_open = pdfplumber.open

    def _slow_open(*args: Any, **kwargs: Any) -> Any:
        time.sleep(0.2)  # stand in for a PDF that is expensive to open
        return real_open(*args, **kwargs)

    # On `pdfplumber` itself, for the reason given in the Word and Excel cases:
    # the connector imports it inside `ingest`.
    monkeypatch.setattr(pdfplumber, "open", _slow_open)
    monkeypatch.setattr(_limits, "MAX_SECONDS", 0.05)

    with pytest.raises(DocumentTooComplexError, match="0 of 3 pages"):
        PdfConnector().ingest(path, source_id="s")


def test_a_refused_workbook_does_not_stay_open(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The budget firing must not leak the workbook's file handle.

    The check used to sit between `load_workbook` and the `try` whose `finally`
    calls `wb.close()`, so the one case the budget exists for was the one case
    that skipped the close — leaving the zip handle to a finaliser, which on
    Windows is long enough to keep the uploaded file locked against the caller
    deleting it.

    Asserted on `close()` actually being called, rather than on the file being
    deletable, because the latter passes on this platform either way and would
    prove nothing.
    """
    import openpyxl as real_openpyxl
    from nlqueries.document_connectors.excel import ExcelConnector

    path = _sheet(tmp_path / "leak.xlsx", rows=10)
    real_load = real_openpyxl.load_workbook
    closed: list[bool] = []

    def _tracking_load(*args: Any, **kwargs: Any) -> Any:
        wb = real_load(*args, **kwargs)
        real_close = wb.close

        def _close() -> None:
            closed.append(True)
            real_close()

        wb.close = _close  # type: ignore[method-assign]
        time.sleep(0.2)
        return wb

    monkeypatch.setattr(real_openpyxl, "load_workbook", _tracking_load)
    monkeypatch.setattr(_limits, "MAX_SECONDS", 0.05)

    with pytest.raises(DocumentTooComplexError):
        ExcelConnector().ingest(path, source_id="s")

    assert closed, "the workbook was not closed when the budget refused it"
