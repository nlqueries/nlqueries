"""Bounds on what document extraction may consume (security-review finding 2.8).

Real files, real parsers. ``tests/test_document_connectors.py`` mocks pdfplumber
and python-docx so it can run without the extras; that is the right trade for
testing chunk shapes and the wrong one here. A limit on what a parser consumes is
only meaningful against the parser -- a mocked ``openpyxl`` would confirm the
arithmetic in ``_limits`` and say nothing about whether a real ``.xlsx`` is
refused. The extras are in the ``dev`` group for this reason.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from nlqueries.document_connectors import _limits
from nlqueries.document_connectors._limits import (
    DocumentTooComplexError,
    check_archive_expansion,
)

openpyxl = pytest.importorskip("openpyxl")
docx = pytest.importorskip("docx")


def _sheet(path: Path, rows: int, cell: str = "value") -> Path:
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("data")
    for _ in range(rows):
        ws.append([cell] * 5)
    wb.save(str(path))
    return path


def _expansion(path: Path) -> tuple[int, int]:
    """(compressed, expanded) totals from the zip central directory."""
    with zipfile.ZipFile(path) as archive:
        info = archive.infolist()
        return sum(i.compress_size for i in info), sum(i.file_size for i in info)


# ---------------------------------------------------------------------------


def test_an_ordinary_spreadsheet_is_not_refused(tmp_path: Path) -> None:
    """The control, and the one that matters most.

    Every limit here is a ceiling over real documents. A guard that refuses an
    ordinary file is worse than no guard: it takes a working feature away, and
    reports the failure as though the user had done something wrong.
    """
    path = _sheet(tmp_path / "ordinary.xlsx", rows=2_000)
    check_archive_expansion(path)  # must not raise

    from nlqueries.document_connectors.excel import ExcelConnector

    assert ExcelConnector().ingest(path, source_id="s"), "ordinary sheet produced no chunks"


def test_a_spreadsheet_over_the_byte_ceiling_is_refused(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Refused on total expanded size, whatever the ratio.

    The ceiling is lowered rather than the file made enormous: the behaviour
    under test is the comparison, and writing 400 MiB to reach it would cost
    minutes per run and exercise the same branch.
    """
    path = _sheet(tmp_path / "big.xlsx", rows=5_000)
    _compressed, expanded = _expansion(path)

    monkeypatch.setattr(_limits, "MAX_EXPANDED_BYTES", expanded - 1)
    monkeypatch.setattr(_limits, "MAX_EXPANSION_RATIO", 10**9)  # isolate the byte rule

    with pytest.raises(DocumentTooComplexError, match="expands to"):
        check_archive_expansion(path)


def test_a_highly_compressible_spreadsheet_is_refused_on_ratio(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The bomb case: small on disk, large in memory.

    The byte ceiling alone does not catch this, because the file is tiny. The
    achieved ratio is asserted first, so a change in compression behaviour shows
    up as a failing premise rather than as a test that silently stops testing
    anything.
    """
    path = _sheet(tmp_path / "compressible.xlsx", rows=20_000, cell="A" * 200)
    compressed, expanded = _expansion(path)
    ratio = expanded / compressed
    assert ratio > 20, f"premise failed: this content only reached {ratio:.1f}x"

    monkeypatch.setattr(_limits, "MAX_EXPANDED_BYTES", 10**12)  # isolate the ratio rule
    monkeypatch.setattr(_limits, "MAX_EXPANSION_RATIO", int(ratio) - 1)

    with pytest.raises(DocumentTooComplexError, match="expands"):
        check_archive_expansion(path)


def test_a_word_document_goes_through_the_same_guard(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """`.docx` is a zip too, and was unguarded for the same reason `.xlsx` was."""
    path = tmp_path / "doc.docx"
    document = docx.Document()
    for _ in range(2_000):
        document.add_paragraph("some text that compresses well " * 5)
    document.save(str(path))

    _compressed, expanded = _expansion(path)
    monkeypatch.setattr(_limits, "MAX_EXPANDED_BYTES", expanded - 1)
    monkeypatch.setattr(_limits, "MAX_EXPANSION_RATIO", 10**9)

    with pytest.raises(DocumentTooComplexError):
        check_archive_expansion(path)


def test_a_non_zip_file_is_left_alone(tmp_path: Path) -> None:
    """A `.pdf` is not an archive, and a corrupt file is the parser's to report.

    Returning quietly rather than raising is deliberate: "not a valid xlsx" from
    openpyxl is a better message than anything this function could invent about
    a file it cannot read.
    """
    plain = tmp_path / "not-an-archive.pdf"
    plain.write_bytes(b"%PDF-1.4 this is not a zip")
    check_archive_expansion(plain)  # must not raise

    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_bytes(b"PK\x03\x04 truncated and broken")
    check_archive_expansion(corrupt)  # must not raise


def test_an_archive_of_stored_entries_is_not_an_expansion(tmp_path: Path) -> None:
    """compress_size == file_size, so the ratio is 1 and the guard must not fire.

    The zero-denominator case lives here too: an empty archive reports both
    totals as zero, and an unguarded ``expanded / compressed`` would raise
    ZeroDivisionError rather than refuse or allow.
    """
    stored = tmp_path / "stored.xlsx"
    with zipfile.ZipFile(stored, "w", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr("payload.bin", b"x" * 100_000)
    check_archive_expansion(stored)  # must not raise

    empty = tmp_path / "empty.xlsx"
    with zipfile.ZipFile(empty, "w"):
        pass
    check_archive_expansion(empty)  # must not raise


def test_the_limits_come_from_the_environment(monkeypatch: pytest.MonkeyPatch) -> None:
    """Configurable, and a malformed value falls back rather than raising.

    These are ceilings; a typo in a deployment's environment should leave the
    documented default in force, not stop every ingest.
    """
    monkeypatch.setenv("NLQ_MAX_DOCUMENT_EXPANDED_BYTES", "123")
    assert _limits._int_env("NLQ_MAX_DOCUMENT_EXPANDED_BYTES", 999) == 123

    for bad in ("not-a-number", "0", "-5", ""):
        monkeypatch.setenv("NLQ_MAX_DOCUMENT_EXPANDED_BYTES", bad)
        assert _limits._int_env("NLQ_MAX_DOCUMENT_EXPANDED_BYTES", 999) == 999, bad

    monkeypatch.delenv("NLQ_MAX_DOCUMENT_EXPANDED_BYTES")
    assert _limits._int_env("NLQ_MAX_DOCUMENT_EXPANDED_BYTES", 999) == 999


def test_the_connector_refuses_before_openpyxl_opens_the_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The wiring, not the function.

    ``check_archive_expansion`` having its own tests says nothing about whether
    the connector calls it, and the connector is the only caller that matters.
    Asserted by making ``load_workbook`` fail the test if it is reached, so a
    guard that ran too late would be caught as well as one that never ran.
    """
    import openpyxl as real_openpyxl
    from nlqueries.document_connectors import excel as excel_module

    path = _sheet(tmp_path / "x.xlsx", rows=1_000)
    _compressed, expanded = _expansion(path)
    monkeypatch.setattr(_limits, "MAX_EXPANDED_BYTES", expanded - 1)
    monkeypatch.setattr(_limits, "MAX_EXPANSION_RATIO", 10**9)

    def _must_not_be_called(*args: object, **kwargs: object) -> None:
        raise AssertionError("openpyxl opened a document the guard should have refused")

    monkeypatch.setattr(real_openpyxl, "load_workbook", _must_not_be_called)

    with pytest.raises(DocumentTooComplexError):
        excel_module.ExcelConnector().ingest(path, source_id="s")
