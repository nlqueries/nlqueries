"""
nlqueries.document_connectors.excel
~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
Microsoft Excel (.xlsx) document connector using ``openpyxl`` for sheet
and row extraction.

Chunking strategy: row-batch-to-text per sheet.  Each sheet is processed
in batches of 50 rows.  The first row of a sheet is tested with a header
heuristic: if all non-empty cells are non-numeric strings (i.e. ``str`` type,
not ``int`` or ``float``), those values are used as column-name prefixes for
every row in the sheet.  Each data row is serialised as
``"col1: val1 | col2: val2 | ..."``.

``page_number`` maps to the sheet index (1-based).

``.xls`` and ``.csv`` are not supported — openpyxl handles ``.xlsx`` only.

Requires the ``docs`` optional dependency group:
    pip install "nlqueries-core[docs]"
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from nlqueries.document_connectors import _limits
from nlqueries.document_connectors._limits import (
    DocumentTooComplexError,
    ExtractionBudget,
    check_archive_expansion,
)
from nlqueries.document_connectors.base import DocumentChunk, DocumentConnector

_BATCH_SIZE = 50


def _make_chunk_id(source_id: str, sheet_index: int, batch_index: int) -> str:
    """Deterministic 16-char hex ID: sha256(source_id:sheet_index:batch_index)[:16]."""
    raw = f"{source_id}:{sheet_index}:{batch_index}"
    return hashlib.sha256(raw.encode()).hexdigest()[:16]


def _cell_str(value: Any) -> str:
    """Convert a cell value to a stripped string; return empty string for None."""
    if value is None:
        return ""
    return str(value).strip()


def _is_header_row(values: list[Any]) -> bool:
    """Heuristic: all non-empty cells are non-numeric strings (``str`` type)."""
    non_empty = [v for v in values if v is not None and str(v).strip() != ""]
    if not non_empty:
        return False
    return all(isinstance(v, str) for v in non_empty)


def _rows_to_text(rows: list[list[Any]], headers: list[str] | None) -> str:
    """Serialise a batch of data rows to a single text block.

    When *headers* is provided each row becomes ``"col1: val1 | col2: val2 | ..."``.
    Without headers the raw cell values are joined with `` | ``.
    """
    lines: list[str] = []
    for row in rows:
        if headers:
            pairs = [f"{headers[i]}: {_cell_str(v)}" for i, v in enumerate(row) if i < len(headers)]
        else:
            pairs = [_cell_str(v) for v in row]
        line = " | ".join(p for p in pairs if p)
        if line:
            lines.append(line)
    return "\n".join(lines)


def _batch_chunk(
    batch: list[list[Any]],
    batch_index: int,
    *,
    source_id: str,
    source_path: Path,
    sheet_title: str,
    sheet_index: int,
    headers: list[str] | None,
    first_data_row_num: int,
) -> DocumentChunk | None:
    """One batch's chunk, or ``None`` when the batch serialises to nothing.

    ``None`` rather than an empty chunk, and the caller advances *batch_index*
    regardless: a wholly blank batch leaves a gap in the emitted indexes. That
    was the behaviour before this connector streamed -- the old code built every
    row of a sheet, enumerated over fixed offsets and ``continue``d on blank text
    -- and the ids are derived from the index, so changing it would renumber
    chunks for any sheet containing a blank stretch.
    """
    text = _rows_to_text(batch, headers)
    if not text.strip():
        return None
    row_start = first_data_row_num + batch_index * _BATCH_SIZE
    row_end = row_start + len(batch) - 1
    return DocumentChunk(
        chunk_id=_make_chunk_id(source_id, sheet_index, batch_index),
        source_id=source_id,
        source_name=source_path.name,
        page_number=sheet_index,
        chunk_index=batch_index,
        text=text,
        metadata={
            "connector": "excel",
            "file_path": str(source_path),
            "sheet_name": sheet_title,
            "row_range": f"{row_start}-{row_end}",
        },
    )


class ExcelConnector(DocumentConnector):
    """Extract and chunk text from Excel (.xlsx) files using openpyxl.

    Each worksheet maps to a ``page_number`` equal to its 1-based sheet index.
    Rows are consumed in batches of ``_BATCH_SIZE`` (50).  The first row of each
    sheet is analysed with ``_is_header_row()``; when all its non-empty cells are
    plain strings, those values are treated as column headers and used to prefix
    every subsequent data row.
    """

    def ingest(self, source: str | Path, source_id: str) -> list[DocumentChunk]:
        try:
            import openpyxl
        except ImportError as exc:
            raise ImportError(
                "openpyxl is required for ExcelConnector. "
                "Install it with: pip install 'nlqueries-core[docs]'"
            ) from exc

        source_path = Path(source)
        # Refuse a document that would expand beyond the limits before the
        # parser opens it. Reads the zip central directory only; nothing is
        # decompressed -- see `_limits` for the measured ratios.
        check_archive_expansion(source_path)

        # Started BEFORE `load_workbook`, matching `pdf.py` and `word.py`.
        # `read_only=True` makes row iteration lazy; it does not make opening
        # free -- the manifest, the styles and (for a file written by Excel) the
        # shared-string table are read eagerly. Measured here at ~0.33s for a
        # 60,000-row workbook before a single row is iterated, which is work the
        # clock should be accountable for.
        #
        # An earlier revision built the budget after this call, so that phase sat
        # outside the clock. That was the same defect review found in `word.py`
        # the round before, and it survived here because I fixed the instance I
        # was shown rather than looking for its siblings.
        budget = ExtractionBudget(name=source_path.name)
        wb = openpyxl.load_workbook(str(source_path), read_only=True, data_only=True)

        chunks: list[DocumentChunk] = []
        rows_seen = 0

        try:
            # Inside the `try`, so `finally: wb.close()` runs when it fires --
            # which is exactly the case the budget exists for. Raising between
            # `load_workbook` and the `try` left the read-only workbook and the
            # zip handle it holds to a finaliser, and on Windows that is long
            # enough to keep the uploaded file locked against the caller
            # deleting it. Checked here rather than earlier changes nothing about
            # *when* it fires: no rows have been read yet either way.
            budget.check("opening the workbook")

            for sheet_index, sheet in enumerate(wb.worksheets, start=1):
                headers: list[str] | None = None
                first_data_row_num = 1
                header_detected = False
                batch: list[list[Any]] = []
                batch_index = 0

                for row_num, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    values = list(row)
                    if not header_detected:
                        header_detected = True
                        if _is_header_row(values):
                            headers = [_cell_str(v) or f"col{i + 1}" for i, v in enumerate(values)]
                            first_data_row_num = row_num + 1
                            continue  # header consumed; not a data row
                    batch.append(values)
                    # Counted only when the row has content. `iter_rows` yields
                    # every materialised row, and whole-column formatting leaves
                    # thousands of empty ones behind -- measured: a 200-row sheet
                    # with such padding yields 5,201 tuples. Counting those would
                    # refuse a workbook whose owner can see 200 rows in it.
                    #
                    # Deliberately separate from the batching, which still takes
                    # every row: blank batches are what produce the index gaps
                    # that `_batch_chunk` preserves.
                    if any(v is not None and str(v).strip() != "" for v in values):
                        rows_seen += 1
                    if rows_seen > _limits.MAX_ROWS:
                        raise DocumentTooComplexError(
                            f"{source_path.name} has more than {_limits.MAX_ROWS} rows; "
                            f"extraction stopped at sheet {sheet.title!r}, row {row_num}."
                        )
                    if len(batch) == _BATCH_SIZE:
                        chunk = _batch_chunk(
                            batch,
                            batch_index,
                            source_id=source_id,
                            source_path=source_path,
                            sheet_title=sheet.title,
                            sheet_index=sheet_index,
                            headers=headers,
                            first_data_row_num=first_data_row_num,
                        )
                        if chunk is not None:
                            chunks.append(chunk)
                        batch_index += 1
                        batch = []
                        budget.check(f"{rows_seen} rows")

                if batch:
                    chunk = _batch_chunk(
                        batch,
                        batch_index,
                        source_id=source_id,
                        source_path=source_path,
                        sheet_title=sheet.title,
                        sheet_index=sheet_index,
                        headers=headers,
                        first_data_row_num=first_data_row_num,
                    )
                    if chunk is not None:
                        chunks.append(chunk)
                budget.check(f"sheet {sheet_index}")
        finally:
            wb.close()

        return chunks

    def supports(self, source: str | Path) -> bool:
        return Path(source).suffix.lower() == ".xlsx"
